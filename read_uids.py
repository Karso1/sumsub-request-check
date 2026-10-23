#!/usr/bin/env python3
"""Read UID values as strings without losing 19-digit precision."""
import json
import sys
from openpyxl import load_workbook

path, sheet_name = sys.argv[1:3]
book = load_workbook(path, read_only=True, data_only=True)
if sheet_name == '__FIRST__':
    sheet_name = book.sheetnames[0]
if sheet_name not in book.sheetnames:
    raise SystemExit(f"Worksheet {sheet_name!r} not found. Available: {book.sheetnames}")
sheet = book[sheet_name]
headers = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))
try:
    column = next(i for i, value in enumerate(headers) if str(value).strip().upper() == "UID")
except StopIteration:
    raise SystemExit("No UID column found")
result = []
seen = set()
for row_number, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), 2):
    value = row[column]
    if value is None or str(value).strip() == "":
        continue
    uid = str(value).strip()
    if not uid.isdecimal():
        raise SystemExit(f"Invalid UID on row {row_number}: {uid!r}")
    if uid not in seen:
        result.append({"uid": uid, "row": row_number})
        seen.add(uid)
print(json.dumps(result, ensure_ascii=False))
